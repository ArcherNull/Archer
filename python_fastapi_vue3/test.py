<<<<<<< HEAD
'''
Author: junsong Chen 779217162@qq.com
Date: 2025-06-03 08:51:05
LastEditTime: 2025-06-03 08:52:15
Description: 
'''
arr=[]
print(arr is not None)
=======
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

# 创建新工作簿
wb = Workbook()
ws = wb.active

# 设置字体样式
font = Font(name='Arial', bold=True, size=14)
ws['A1'].font = font

# 设置背景颜色
fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
ws['B1'].fill = fill

# 设置边框
border = Border(left=Side(style='thin'), right=Side(style='thin'),
top=Side(style='thin'), bottom=Side(style='thin'))
ws['C1'].border = border

# 添加数据到工作表
ws['A1'] = 'Name'
ws['B1'] = 'Age'
ws['C1'] = 'City'
ws.append(['John', 28, 'New York'])
ws.append(['Anna', 34, 'Paris'])
ws.append(['Peter', 29, 'Berlin'])
ws.append(['Linda', 32, 'London'])
 

# 保存工作簿
wb.save('styled.xlsx')
>>>>>>> 9b94e5dae18a75126b72fc9b427907e32381e67c
